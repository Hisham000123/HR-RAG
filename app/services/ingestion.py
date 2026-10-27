from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument  # pip install python-docx
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader  # PDF needs: pip install pypdf


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif suffix in {".txt", ".md"}:
        docs = TextLoader(str(path), encoding="utf-8", autodetect_encoding=True).load()
    elif suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        docs = [Document(page_content=text, metadata={"source": str(path)})]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    # Drop empty pages (e.g. scanned PDFs with no extractable text)
    docs = [d for d in docs if d.page_content and d.page_content.strip()]
    if not docs:
        raise ValueError(
            "No text could be extracted. If this is a scanned PDF, it needs OCR first."
        )
    return docs


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900, chunk_overlap=120, add_start_index=True
    )
    return splitter.split_documents(list(docs))